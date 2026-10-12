
import re
import sys
import json
from pathlib import Path
from openpyxl import load_workbook

excel_path = Path(
    sys.argv[1] if len(sys.argv) > 1
    else "题库搜索记录-题目导出.xlsx"
)
cache_path = Path("cache.json")


def normalize_question(question):
    """与 SuperStar 原项目的题目标题处理规则保持一致。"""
    question = str(question).strip()
    question = re.sub(r"^\d+", "", question)
    question = re.sub(r"（\d+\.\d+分）$", "", question)
    return question


# 读取已有缓存，避免覆盖历史数据
if cache_path.exists():
    try:
        cache = json.loads(cache_path.read_text(encoding="utf-8"))
        if not isinstance(cache, dict):
            cache = {}
    except (json.JSONDecodeError, OSError):
        cache = {}
else:
    cache = {}

before = len(cache)

if not excel_path.exists():
    raise FileNotFoundError(f"找不到 Excel 文件：{excel_path}")

workbook = load_workbook(excel_path, read_only=True, data_only=True)
sheet = workbook.active

headers = [
    str(cell).strip() if cell is not None else ""
    for cell in next(sheet.iter_rows(values_only=True))
]

try:
    question_col = headers.index("查询题目")
    answer_col = headers.index("搜到答案")
except ValueError as exc:
    raise ValueError(
        "Excel 必须包含“查询题目”和“搜到答案”两列"
    ) from exc

added = 0
skipped = 0

for row in sheet.iter_rows(min_row=2, values_only=True):
    if len(row) <= max(question_col, answer_col):
        continue

    question = row[question_col]
    answer = row[answer_col]

    if question is None or answer is None:
        skipped += 1
        continue

    question = normalize_question(question)
    answer = str(answer).strip()

    if not question or not answer:
        skipped += 1
        continue

    # 不覆盖已有答案，优先保留原有缓存
    if question not in cache:
        cache[question] = answer
        added += 1
    else:
        skipped += 1

workbook.close()

cache_path.write_text(
    json.dumps(cache, ensure_ascii=False, indent=2),
    encoding="utf-8"
)

print(f"导入完成：新增 {added} 道题")
print(f"跳过记录：{skipped} 条")
print(f"缓存总数：{len(cache)} 道")
print(f"原缓存数量：{before} 道")
